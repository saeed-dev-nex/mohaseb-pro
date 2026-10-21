"""Characterization tests for report/exports/dashboard/chat.

These lock the exact numbers and shapes the frontend depends on, so the
stage-3 refactor (service layer, openpyxl, views package) cannot quietly
change behaviour.
"""

import io
from decimal import Decimal

from django.contrib.auth.models import User
from django.core.cache import cache
from django.utils import timezone
from rest_framework.test import APITestCase

from api.models import (
    Customer, Expense, ExpenseCategory, Invoice, InvoiceItem, InvoicePayment, Service,
)


class ReportSeed(APITestCase):
    def setUp(self):
        cache.clear()
        self.user = User.objects.create_user(username='reporter', password='test-pass-123')
        token = self.client.post('/api/token/', {'username': 'reporter', 'password': 'test-pass-123'}, format='json')
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token.data['access']}")

        self.ali = Customer.objects.create(name='علی رضایی')
        self.sara = Customer.objects.create(name='سارا محمدی')
        self.service = Service.objects.create(title='خلافی')
        now = timezone.now()

        # final = 100000 + 10% tax - 5000 discount + 7000 commission = 112000
        self.inv_paid = Invoice.objects.create(
            customer=self.ali, status='paid', date=now, total_amount=Decimal('100000'),
            tax=Decimal('10'), discount=Decimal('5000'), commission_total=Decimal('7000'))
        InvoicePayment.objects.create(invoice=self.inv_paid, amount=Decimal('50000'), method='cash', created_by=self.user)
        InvoicePayment.objects.create(invoice=self.inv_paid, amount=Decimal('62000'), method='card', created_by=self.user)
        InvoiceItem.objects.create(invoice=self.inv_paid, service=self.service, quantity=2,
                                     main_price=Decimal('10000'), extra_price=Decimal('5000'))

        self.inv_pending = Invoice.objects.create(
            customer=self.ali, status='pending', date=now, total_amount=Decimal('50000'))
        InvoicePayment.objects.create(invoice=self.inv_pending, amount=Decimal('20000'), method='transfer', created_by=self.user)
        InvoiceItem.objects.create(invoice=self.inv_pending, service_title_snapshot='پیگیری', quantity=1,
                                     main_price=Decimal('5000'))

        self.inv_cancelled = Invoice.objects.create(
            customer=self.sara, status='cancelled', date=now, total_amount=Decimal('30000'))

        category = ExpenseCategory.objects.create(name='خوراک')
        Expense.objects.create(title='ناهار', category=category, amount=Decimal('8000'), date=now.date())
        Expense.objects.create(title='سایر', amount=Decimal('2000'), date=now.date())


class ReportTests(ReportSeed):
    def test_report_totals(self):
        data = self.client.get('/api/invoices/report/').data
        self.assertEqual(data['total_revenue'], 132000)
        self.assertEqual(data['total_commission'], 7000)
        self.assertEqual(data['pending_amount'], 30000)
        self.assertEqual(data['total_count'], 3)
        self.assertEqual(data['invoice_amount'], 162000)
        self.assertEqual(data['expense_total'], 10000)
        self.assertEqual(data['net_profit'], -3000)

    def test_report_settled_after_partial_payments_counts_full_amount(self):
        # Regression: revenue follows the serializer rule — an invoice
        # settled after staged payments counts at its full final amount,
        # while income (commission) still counts commission only.
        from django.utils import timezone
        inv = Invoice.objects.create(
            customer=self.ali, status='paid', date=timezone.now(),
            total_amount=Decimal('40000'), commission_total=Decimal('3000'))
        InvoicePayment.objects.create(invoice=inv, amount=Decimal('10000'), created_by=self.user)
        data = self.client.get('/api/invoices/report/').data
        # final = 40000 total + 3000 commission = 43000, counted in full.
        self.assertEqual(data['total_revenue'], 132000 + 43000)
        self.assertEqual(data['total_commission'], 7000 + 3000)

    def test_report_breakdowns(self):
        data = self.client.get('/api/invoices/report/').data
        self.assertEqual(data['by_status']['paid'], {'label': 'پرداخت شده', 'count': 1, 'amount': 112000})
        self.assertEqual(data['by_status']['pending']['amount'], 30000)
        self.assertEqual(data['by_status']['cancelled']['amount'], 30000)
        self.assertEqual(data['monthly'][0]['count'], 3)
        self.assertEqual(data['monthly'][0]['amount'], 192000)
        self.assertEqual(data['monthly'][0]['received'], 132000)
        self.assertEqual(data['top_customers'][0], {'name': 'علی رضایی', 'count': 2, 'amount': 162000})
        services = {row['name']: row['amount'] for row in data['top_services']}
        self.assertEqual(services['خلافی'], 30000)
        self.assertEqual(services['پیگیری'], 5000)
        methods = {row['name']: row['amount'] for row in data['payment_methods']}
        self.assertEqual(methods, {'نقدی': 50000, 'کارت‌خوان': 62000, 'کارت‌به‌کارت / حواله': 20000})
        categories = {row['name']: row['amount'] for row in data['expense_categories']}
        self.assertEqual(categories, {'خوراک': 8000, 'بدون دسته': 2000})
        self.assertEqual(len(data['invoices']), 3)

    def test_report_status_filter(self):
        data = self.client.get('/api/invoices/report/?status=pending').data
        self.assertEqual(data['total_count'], 1)
        self.assertEqual(data['total_revenue'], 20000)


class ExportTests(ReportSeed):
    def test_excel_workbook_contents(self):
        from openpyxl import load_workbook

        response = self.client.get('/api/invoices/export_excel/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['خلاصه', 'فاکتورها', 'هزینه‌ها', 'پرداخت‌ها'])
        summary = {row[0].value: row[1].value for row in workbook['خلاصه'].iter_rows() if row[0].value}
        self.assertEqual(summary['تعداد فاکتور'], 3)
        self.assertEqual(summary['مبلغ دریافت‌شده'], 132000)
        self.assertEqual(summary['کارمزد فاکتورهای پرداخت‌شده'], 7000)
        self.assertEqual(summary['مبلغ هزینه‌ها'], 10000)
        self.assertEqual(summary['سود خالص'], -3000)
        invoices_sheet = workbook['فاکتورها']
        self.assertEqual(invoices_sheet.max_row, 4)  # header + 3 invoices
        numbers = {row[0].value for row in invoices_sheet.iter_rows(min_row=2)}
        self.assertIn(self.inv_paid.invoice_number, numbers)
        payments_sheet = workbook['پرداخت‌ها']
        self.assertEqual(payments_sheet.max_row, 4)  # header + 3 payments

    def test_csv_export(self):
        response = self.client.get('/api/invoices/export_csv/')
        self.assertEqual(response.status_code, 200)
        text = response.content.decode('utf-8-sig')
        self.assertIn('شماره فاکتور', text)
        self.assertIn(self.inv_paid.invoice_number, text)
        self.assertIn(self.inv_cancelled.invoice_number, text)


class DashboardAndChatTests(ReportSeed):
    def test_dashboard_stats(self):
        response = self.client.get('/api/dashboard/stats/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['total_customers'], 2)
        self.assertEqual(response.data['total_invoices'], 3)

    def test_chat_public_room_flow(self):
        other = User.objects.create_user(username='chatter', password='test-pass-123')
        response = self.client.post('/api/chat/', {'text': 'سلام'}, format='json')
        self.assertEqual(response.status_code, 201, response.content)
        listing = self.client.get('/api/chat/')
        self.assertEqual(listing.status_code, 200)
        self.assertEqual(len(listing.data), 1)
        self.assertEqual(listing.data[0]['text'], 'سلام')
        users = self.client.get('/api/chat/users/')
        self.assertEqual(users.status_code, 200)
        self.assertEqual([row['username'] for row in users.data], ['chatter'])
        unread = self.client.get('/api/chat/unread/')
        self.assertEqual(unread.data['count'], 0)  # own messages do not notify the sender
