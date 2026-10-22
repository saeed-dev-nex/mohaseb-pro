"""Financial boundary regressions; synthetic data in Django's isolated test DB."""

import csv
import io
from datetime import datetime, timezone as datetime_timezone
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.test import override_settings
from django.utils import timezone
from rest_framework.test import APITestCase

from api.models import (
    Customer, Invoice, InvoiceItem, InvoiceNumberCounter, InvoicePayment,
    Expense, Staff, Cheque, Service, InvoiceSettings, UserModuleAccess,
)
from api.serializers import (
    InvoiceSerializer, CustomerListSerializer, CustomerDetailSerializer,
    InvoicePaymentSerializer, ExpenseSerializer, StaffSerializer, ChequeSerializer,
    ServiceSerializer, UserSerializer,
)
from api.services.money import payment_totals
from api.services.reports import build_report_data, build_excel_bytes


class FinancialRegressionTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='finance-manager', password='fixture', is_staff=True)
        self.client.force_authenticate(self.user)

    def payload(self, **overrides):
        payload = {'items': [{'service_title_snapshot': 'خدمت تست', 'quantity': 2, 'main_price': '1000'}]}
        payload.update(overrides)
        return payload

    def quick_payload(self, **overrides):
        return self.payload(**{'customer_name': 'مشتری آزمایشی', 'customer_phone': '09120000000', **overrides})

    def test_invalid_invoice_inputs_rejected_before_any_customer_or_invoice_is_created(self):
        invalid = [
            {'items': []}, {'items': 'invalid'}, {'items': [{'quantity': 0}]},
            {'items': [{'quantity': -1}]}, {'items': [{'quantity': '1.5'}]},
            {'items': [{'main_price': '-1'}]}, {'items': [{'extra_price': '-1'}]},
            {'items': [{'commission': '-1'}]}, {'items': [{'service': 999999}]},
            {'tax': '-0.01'}, {'tax': '100.01'}, {'tax': 'NaN'}, {'tax': 'Infinity'},
            {'discount': '-1'}, {'commission_total': '-1'}, {'total_amount': '-1'},
            {'discount': '2001'}, {'status': 'made-up'}, {'card_payment_source': 'made-up'},
            {'items': [{'quantity': 1001, 'main_price': '999999999999'}]},
            {'items': [{'quantity': 600, 'main_price': '999999999999'}], 'tax': '100'},
        ]
        for changes in invalid:
            for endpoint, payload in [('/api/invoices/', self.payload(**changes)),
                                      ('/api/invoices/quick_create/', self.quick_payload(**changes))]:
                with self.subTest(changes=changes, endpoint=endpoint):
                    response = self.client.post(endpoint, payload, format='json')
                    self.assertEqual(response.status_code, 400, response.data)
                    self.assertFalse(Customer.objects.exists())
                    self.assertFalse(Invoice.objects.exists())
                    self.assertFalse(InvoiceNumberCounter.objects.exists())

    def test_quick_create_missing_items_or_invalid_customer_is_rejected(self):
        for payload in [{'customer_name': 'test'}, self.quick_payload(customer_name=''),
                        self.quick_payload(customer_name='x' * 256), self.quick_payload(customer_phone='1' * 51)]:
            response = self.client.post('/api/invoices/quick_create/', payload, format='json')
            self.assertEqual(response.status_code, 400, response.data)
            self.assertFalse(Customer.objects.exists())

    def test_quick_and_normal_creation_share_calculation_and_optional_fields(self):
        InvoiceSettings.objects.create(office_card_fee_threshold=100)
        payload = self.quick_payload(tax='9.25', discount='25', commission_total='100', card_payment_source='office',
                                     tags=['synthetic'], vehicle_year='1405', custom_fields={'test': 'value'},
                                     green_card_owner_name='تست', status='paid', paper_size='a4')
        response = self.client.post('/api/invoices/quick_create/', payload, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(set(response.data), {'invoice', 'customer', 'message'})
        invoice = Invoice.objects.get(pk=response.data['invoice']['id'])
        # subtotal 2000 + fee 20; 2020 + 9.25% - 25 + 100 = 2281.85.
        self.assertEqual(invoice.total_amount, Decimal('2020'))
        self.assertEqual(invoice.final_amount, Decimal('2281.85'))
        self.assertEqual(invoice.tags, ['synthetic'])
        self.assertEqual(invoice.vehicle_year, '1405')
        self.assertEqual(invoice.custom_fields, {'test': 'value'})
        self.assertEqual(invoice.paid_by, self.user)
        self.assertEqual(invoice.items.filter(is_office_card_fee=True).get().main_price, 20)
        self.assertTrue(response.data['customer']['is_new'])
        normal = self.client.post('/api/invoices/', {**payload, 'customer': invoice.customer_id}, format='json')
        self.assertEqual(normal.status_code, 201, normal.data)
        self.assertEqual(normal.data['final_amount'], response.data['invoice']['final_amount'])
        self.assertEqual(normal.data['total_amount'], response.data['invoice']['total_amount'])

    def test_generated_fee_must_fit_its_item_field(self):
        InvoiceSettings.objects.create(office_card_fee_threshold=1)
        response = self.client.post('/api/invoices/quick_create/', self.quick_payload(
            card_payment_source='office', items=[{'quantity': 200, 'main_price': '999999999999'}]), format='json')
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Customer.objects.exists())

    def test_quick_creation_rolls_back_customer_invoice_items_and_number_on_late_failure(self):
        with patch.object(InvoiceSerializer, '_sync_office_card_fee', side_effect=RuntimeError('synthetic failure')):
            with self.assertRaises(RuntimeError):
                self.client.post('/api/invoices/quick_create/', self.quick_payload(), format='json')
        self.assertFalse(Customer.objects.exists())
        self.assertFalse(Invoice.objects.exists())
        self.assertFalse(InvoiceItem.objects.exists())
        self.assertFalse(InvoiceNumberCounter.objects.exists())

    def test_normal_creation_rolls_back_on_audit_failure(self):
        with patch('api.views.invoices.record_audit', side_effect=RuntimeError('synthetic failure')):
            with self.assertRaises(RuntimeError):
                self.client.post('/api/invoices/', self.payload(), format='json')
        self.assertFalse(Invoice.objects.exists())
        self.assertFalse(InvoiceItem.objects.exists())
        self.assertFalse(InvoiceNumberCounter.objects.exists())

    def test_serializer_update_is_atomic_and_preserves_original_rows_on_failure(self):
        invoice = Invoice.objects.create(total_amount=1000, description='original')
        original = InvoiceItem.objects.create(invoice=invoice, main_price=1000)
        serializer = InvoiceSerializer(invoice, data=self.payload(description='changed'), partial=True)
        self.assertTrue(serializer.is_valid(), serializer.errors)
        with patch.object(InvoiceSerializer, '_sync_office_card_fee', side_effect=RuntimeError('synthetic failure')):
            with self.assertRaises(RuntimeError):
                serializer.save()
        invoice.refresh_from_db()
        self.assertEqual(invoice.total_amount, 1000)
        self.assertEqual(invoice.description, 'original')
        self.assertEqual(list(invoice.items.values_list('id', flat=True)), [original.pk])

    def test_partial_update_validates_resulting_final_amount(self):
        invoice = Invoice.objects.create(total_amount=1000, discount=500)
        InvoiceItem.objects.create(invoice=invoice, main_price=1000)
        for changes in [{'discount': 1001}, {'items': [{'main_price': 100}]}, {'tax': 101}, {'status': 'invalid'}]:
            response = self.client.patch(f'/api/invoices/{invoice.pk}/', changes, format='json')
            self.assertEqual(response.status_code, 400, response.data)
        invoice.refresh_from_db()
        self.assertEqual(invoice.final_amount, 500)

    def test_positive_amount_and_percentage_validation_in_serializers_and_models(self):
        customer = Customer.objects.create(name='تست')
        invoice = Invoice.objects.create(customer=customer)
        cases = [
            (InvoicePaymentSerializer, {'amount': 0}), (InvoicePaymentSerializer, {'amount': -1}),
            (InvoicePaymentSerializer, {'amount': '1.5'}),
            (ExpenseSerializer, {'title': 'test', 'amount': 0}),
            (ChequeSerializer, {'bank_name': 'test', 'cheque_number': 'test', 'amount': -1, 'due_date': '2026-10-22'}),
            (StaffSerializer, {'name': 'test', 'commission_rate': -1}),
            (StaffSerializer, {'name': 'test', 'commission_rate': '100.01'}),
            (ServiceSerializer, {'title': 'test', 'default_main_price': -1}),
        ]
        for serializer_class, payload in cases:
            with self.subTest(serializer=serializer_class.__name__, payload=payload):
                serializer = serializer_class(data=payload)
                self.assertFalse(serializer.is_valid())
        invalid_models = [InvoiceItem(invoice=invoice, quantity=0), InvoiceItem(invoice=invoice, main_price=-1),
                          InvoicePayment(invoice=invoice, amount=0), Expense(title='test', amount=-1),
                          Cheque(customer=customer, cheque_number='test', bank_name='test', amount=0,
                                 due_date=timezone.localdate()), Staff(name='test', commission_rate=101),
                          Invoice(total_amount=1, discount=2), Invoice(tax=101), Service(title='test', default_extra_price=-1)]
        for instance in invalid_models:
            with self.subTest(model=type(instance).__name__):
                with self.assertRaises(ValidationError):
                    instance.full_clean()

    def test_payment_endpoint_rejects_fractional_and_oversized_amount_without_writes(self):
        invoice = Invoice.objects.create(total_amount=1000)
        for amount in ['0', '-1', '1.5', '1000000000000000', 'NaN']:
            response = self.client.post(f'/api/invoices/{invoice.pk}/payments/', {'amount': amount}, format='json')
            self.assertEqual(response.status_code, 400, response.data)
        self.assertFalse(InvoicePayment.objects.exists())

    def test_password_validation_on_create_and_update_keeps_blank_update_contract(self):
        for data, instance in [({'username': 'new-user', 'password': '123'}, None),
                               ({'password': '123'}, self.user)]:
            serializer = UserSerializer(instance, data=data, partial=instance is not None)
            self.assertFalse(serializer.is_valid())
            self.assertIn('password', serializer.errors)
        serializer = UserSerializer(self.user, data={'password': ''}, partial=True)
        self.assertTrue(serializer.is_valid(), serializer.errors)
        serializer.save()
        self.assertTrue(self.user.check_password('fixture'))
        serializer = UserSerializer(data={'username': 'new-user', 'password': 'Good!Synthetic-Pass42'})
        self.assertTrue(serializer.is_valid(), serializer.errors)
        self.assertTrue(serializer.save().check_password('Good!Synthetic-Pass42'))
        self.assertEqual(serializer.data['role'], 'operator')
        self.assertEqual(UserSerializer(self.user).data['role'], 'manager')

    def test_legacy_settlement_and_remaining_are_identical_in_all_surfaces(self):
        from openpyxl import load_workbook
        customer = Customer.objects.create(name='تست')
        settled = Invoice.objects.create(customer=customer, status='paid', total_amount=50000)
        pending = Invoice.objects.create(customer=customer, status='pending', total_amount=30000)
        InvoicePayment.objects.create(invoice=settled, amount=20000)
        InvoicePayment.objects.create(invoice=pending, amount=10000)
        self.assertEqual(customer.total_paid, 60000)
        self.assertEqual(customer.total_pending, 20000)
        for serializer_class in [CustomerListSerializer, CustomerDetailSerializer]:
            data = serializer_class(customer).data
            self.assertEqual(data['total_paid'], 60000)
            self.assertEqual(data['total_pending'], 20000)
        report = self.client.get('/api/invoices/report/').data
        self.assertEqual(report['total_revenue'], 60000)
        self.assertEqual(report['pending_amount'], 20000)
        for row in report['invoices']:
            self.assertEqual((row['paid_amount'], row['remaining_amount']),
                             (50000, 0) if row['id'] == settled.pk else (10000, 20000))
        workbook = load_workbook(io.BytesIO(build_excel_bytes()))
        summary = {row[0].value: row[1].value for row in workbook['خلاصه'].iter_rows()}
        self.assertEqual(summary['مبلغ دریافت‌شده'], 60000)
        for row in workbook['فاکتورها'].iter_rows(min_row=2, values_only=True):
            self.assertEqual(row[7:9], (50000, 0) if row[0] == settled.invoice_number else (10000, 20000))
        self.assertEqual(payment_totals(100, 'paid', 120), (Decimal(120), Decimal(0)))

    def test_payment_calculation_uses_prefetch_without_per_invoice_queries(self):
        customer = Customer.objects.create(name='test')
        for _ in range(4):
            invoice = Invoice.objects.create(customer=customer, total_amount=100, status='paid')
            InvoicePayment.objects.create(invoice=invoice, amount=20)
        customer = Customer.objects.prefetch_related('invoice_set__payments', 'documents').get(pk=customer.pk)
        invoices = list(Invoice.objects.select_related('customer', 'created_by', 'paid_by').prefetch_related(
            'payments__created_by', 'items__service', 'edit_requests'))
        with self.assertNumQueries(0):
            self.assertEqual(customer.total_paid, 400)
            self.assertEqual(CustomerDetailSerializer(customer).data['total_paid'], 400)
            self.assertEqual([row['paid_amount'] for row in InvoiceSerializer(invoices, many=True).data], [100] * 4)
        with self.assertNumQueries(3):  # customers, invoices, payments regardless of invoice count
            self.assertEqual(CustomerListSerializer(Customer.objects.all(), many=True).data[0]['total_paid'], 400)

    def test_fractional_tax_is_kept_decimal_until_the_integer_response_boundary(self):
        from openpyxl import load_workbook
        customer = Customer.objects.create(name='fractional-tax')
        for _ in range(2):
            settled = Invoice.objects.create(customer=customer, status='paid', total_amount=100, tax='0.80')
            InvoicePayment.objects.create(invoice=settled, amount=20)
            Invoice.objects.create(customer=customer, status='pending', total_amount=100, tax='0.80')
        self.assertEqual(customer.total_paid, Decimal('201.60'))
        self.assertEqual(customer.total_pending, Decimal('201.60'))
        # Preserve the existing integer response contract, but truncate only
        # after summing Decimal amounts, identically in each summary surface.
        data = CustomerDetailSerializer(customer).data
        report, _ = build_report_data()
        workbook = load_workbook(io.BytesIO(build_excel_bytes()))
        summary = {row[0].value: row[1].value for row in workbook['خلاصه'].iter_rows()}
        self.assertEqual(data['total_paid'], 201)
        self.assertEqual(data['total_pending'], 201)
        self.assertEqual(report['total_revenue'], 201)
        self.assertEqual(report['pending_amount'], 201)
        self.assertEqual(report['by_status']['pending']['amount'], 201)
        self.assertEqual(summary['مبلغ دریافت‌شده'], 201)

    def test_report_export_history_require_reports_access(self):
        invoice = Invoice.objects.create()
        UserModuleAccess.objects.create(user=self.user, reports=False)
        self.user.refresh_from_db()
        for path in ['report', 'export_excel', 'export_csv', f'{invoice.pk}/history']:
            response = self.client.get(f'/api/invoices/{path}/')
            self.assertEqual(response.status_code, 403, path)

    def test_expenses_denied_rejects_combined_reports_and_excel_but_allows_invoice_csv(self):
        UserModuleAccess.objects.create(user=self.user, reports=True, expenses=False)
        self.user.refresh_from_db()
        with patch('api.services.reports.build_report_data') as report, patch('api.services.reports.build_excel_bytes') as excel:
            for path in ['report', 'export_excel']:
                response = self.client.get(f'/api/invoices/{path}/')
                self.assertEqual(response.status_code, 403)
                self.assertIn('هزینه', str(response.data))
            report.assert_not_called()
            excel.assert_not_called()
        self.assertEqual(self.client.get('/api/invoices/export_csv/').status_code, 200)

    def test_invoice_module_gate_remains_required_for_reports(self):
        UserModuleAccess.objects.create(user=self.user, reports=True, invoices=False)
        self.user.refresh_from_db()
        self.assertEqual(self.client.get('/api/invoices/report/').status_code, 403)

    def test_excel_and_csv_neutralize_untrusted_formula_strings(self):
        from openpyxl import load_workbook
        customer = Customer.objects.create(name='=HYPERLINK("https://invalid")')
        self.user.first_name = '+Synthetic'
        self.user.save()
        invoice = Invoice.objects.create(customer=customer, total_amount=100, created_by=self.user, paid_by=self.user,
                                         invoice_number='@untrusted')
        InvoicePayment.objects.create(invoice=invoice, amount=20, created_by=self.user, note='\t=malicious')
        Expense.objects.create(title='-untrusted', description='  =untrusted', amount=1, date=timezone.localdate())
        workbook = load_workbook(io.BytesIO(build_excel_bytes()))
        invoice_row = list(workbook['فاکتورها'].iter_rows(min_row=2))[0]
        self.assertEqual(invoice_row[0].value, "'@untrusted")
        self.assertTrue(invoice_row[1].value.startswith("'="))
        self.assertEqual(invoice_row[9].value, "'+Synthetic")
        self.assertEqual(workbook['پرداخت‌ها']['G2'].value, "'\t=malicious")
        self.assertEqual(workbook['هزینه‌ها']['A2'].value, "'-untrusted")
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, 'f')
        response = self.client.get('/api/invoices/export_csv/')
        self.assertEqual(response.content.count(b'\xef\xbb\xbf'), 1)
        row = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))[1]
        self.assertEqual(row[:2], ["'@untrusted", "'=" + customer.name[1:]])

    @override_settings(TIME_ZONE='Asia/Tehran')
    def test_report_and_exports_use_local_time_at_midnight_boundary(self):
        from openpyxl import load_workbook
        invoice = Invoice.objects.create(date=datetime(2026, 12, 31, 21, 30, tzinfo=datetime_timezone.utc), total_amount=100)
        InvoicePayment.objects.create(invoice=invoice, amount=10, paid_at=invoice.date)
        with timezone.override('Asia/Tehran'):
            data, _ = build_report_data()
            self.assertEqual(data['monthly'][0]['month'], '2027-01')
            workbook = load_workbook(io.BytesIO(build_excel_bytes()))
            self.assertEqual(workbook['فاکتورها']['C2'].value, '2027-01-01')
            self.assertEqual(workbook['پرداخت‌ها']['C2'].value, '2027-01-01 01:00')
            csv_export = self.client.get('/api/invoices/export_csv/').content.decode('utf-8-sig')
            self.assertIn('2027-01-01', csv_export)
